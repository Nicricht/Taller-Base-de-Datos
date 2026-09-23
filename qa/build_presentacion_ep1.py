#!/usr/bin/env python3
"""Genera la presentación ejecutiva EP1 de EDUBIO360.

Diseñada para la Situación Evaluativa 2 de BDY1103 y para ser validada con
qa/validate_presentacion_ep1.py. Usa únicamente python-pptx y evidencias ya
versionadas en el repositorio.
"""

from __future__ import annotations

import zipfile
from datetime import datetime
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.util import Inches, Pt

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "docs/presentacion/EDUBIO360_EP1_PRESENTACION_FINAL.pptx"
EVID = ROOT / "docs/evidencias/ejecucion-oracle"

W, H = Inches(13.333), Inches(7.5)
NAVY = RGBColor(18, 58, 90)
BLUE = RGBColor(46, 92, 161)
GREEN = RGBColor(78, 139, 87)
GOLD = RGBColor(198, 106, 26)
TEXT = RGBColor(31, 41, 55)
MUTED = RGBColor(91, 101, 115)
WHITE = RGBColor(255, 255, 255)
BG = RGBColor(248, 250, 252)
BORDER = RGBColor(214, 222, 232)


def set_bg(slide, color=BG):
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = color


def add_text(slide, text, x, y, w, h, size=18, bold=False, color=TEXT,
             font="Aptos", align=PP_ALIGN.LEFT, valign=MSO_ANCHOR.TOP):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.vertical_anchor = valign
    p = tf.paragraphs[0]
    p.alignment = align
    r = p.add_run()
    r.text = text
    r.font.name = font
    r.font.size = Pt(size)
    r.font.bold = bold
    r.font.color.rgb = color
    return box


def add_title(slide, title, subtitle=None, number=None):
    add_text(slide, title, 0.55, 0.34, 10.9, 0.56, 25, True, NAVY, "Aptos Display")
    if subtitle:
        add_text(slide, subtitle, 0.57, 0.93, 11.3, 0.35, 11.5, False, MUTED)
    if number is not None:
        add_text(slide, str(number), 12.15, 7.0, 0.55, 0.25, 9, False, MUTED, align=PP_ALIGN.RIGHT)


def card(slide, x, y, w, h, title, body, accent=BLUE, body_size=13.5):
    sh = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    sh.fill.solid()
    sh.fill.fore_color.rgb = WHITE
    sh.line.color.rgb = BORDER
    bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(0.08), Inches(h))
    bar.fill.solid()
    bar.fill.fore_color.rgb = accent
    bar.line.fill.background()
    add_text(slide, title, x + 0.22, y + 0.18, w - 0.38, 0.35, 14.5, True, accent)
    add_text(slide, body, x + 0.22, y + 0.62, w - 0.38, h - 0.77, body_size, False, TEXT)


def pill(slide, text, x, y, w, color=BLUE):
    sh = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(0.38))
    sh.fill.solid()
    sh.fill.fore_color.rgb = color
    sh.line.fill.background()
    add_text(slide, text, x, y + 0.01, w, 0.28, 9.5, True, WHITE, align=PP_ALIGN.CENTER, valign=MSO_ANCHOR.MIDDLE)


def add_notes(slide, text):
    tf = slide.notes_slide.notes_text_frame
    tf.clear()
    tf.text = text


def add_picture_if_exists(slide, path: Path, x, y, w, h):
    if not path.exists():
        return False
    slide.shapes.add_picture(str(path), Inches(x), Inches(y), Inches(w), Inches(h))
    return True


def slide1(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    set_bg(s, WHITE)
    add_text(s, "EDUBIO360", 0.75, 1.15, 8.3, 0.8, 34, True, NAVY, "Aptos Display")
    add_text(s, "Evaluación Parcial N° 1", 0.78, 2.0, 7.2, 0.5, 22, True, BLUE)
    add_text(s, "Taller de Base de Datos · BDY1103", 0.8, 2.62, 6.7, 0.35, 14, False, MUTED)
    add_text(s, "Javier Pérez  ·  Nicolás Vega", 0.8, 3.25, 6.8, 0.38, 15.5, False, TEXT)
    add_text(s, "Base de datos Oracle + PL/SQL", 0.8, 3.72, 6.8, 0.38, 14, True, GREEN)
    pill(s, "Presentación individual · 60%", 0.8, 4.35, 2.45, BLUE)
    add_text(s, "Objetivo: explicar y justificar las decisiones técnicas utilizadas para procesar y generar información relevante para EDUBIO360.", 0.8, 5.0, 8.7, 0.85, 16, False, TEXT)
    for i, c in enumerate((NAVY, BLUE, GREEN, GOLD)):
        sh = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(10.1 + i * 0.48), Inches(1.25 + i * 0.56), Inches(2.05), Inches(0.95))
        sh.fill.solid()
        sh.fill.fore_color.rgb = c
        sh.line.fill.background()
    add_notes(s, "Buenos días. Nuestro proyecto se llama EDUBIO360. En esta presentación voy a explicar el problema que resuelve, cómo organizamos los datos y por qué utilizamos RECORD, VARRAY, cursores, loops, excepciones y, a futuro, procedimientos, funciones, packages y triggers.")


def slide2(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    set_bg(s)
    add_title(s, "1. Problema, datos y objetivo", "La información existe, pero en la fuente original está concentrada y repetida.", 2)
    card(s, 0.55, 1.45, 3.75, 2.15, "Problema", "El archivo original contiene instituciones, carreras, comunas, modalidades, jornadas, matrícula y arancel en una estructura plana. Los mismos valores se repiten muchas veces.", GOLD, 12.7)
    card(s, 4.52, 1.45, 3.75, 2.15, "Datos procesados", "106.555 registros\n28 columnas\nInstituciones · carreras · ubicación\nPlanes · costos · vías de ingreso\nMatrícula histórica", BLUE, 13)
    card(s, 8.49, 1.45, 4.25, 2.15, "Objetivo de EDUBIO360", "Organizar la información para consultar y comparar ofertas académicas de forma clara: qué se estudia, dónde, en qué modalidad y cuánto cuesta.", GREEN, 12.7)
    labels = [("Fuente original", "Excel / CSV"), ("Carga temporal", "STAGING_MATRICULA"), ("Normalización", "19 tablas de negocio"), ("PL/SQL", "Procesamiento"), ("Consulta", "Información útil")]
    x = 0.72
    for i, (a, b) in enumerate(labels):
        add_text(s, a, x, 4.45, 2.05, 0.3, 10.5, True, MUTED, align=PP_ALIGN.CENTER)
        sh = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(4.8), Inches(2.05), Inches(0.72))
        sh.fill.solid()
        sh.fill.fore_color.rgb = WHITE
        sh.line.color.rgb = BORDER
        add_text(s, b, x + 0.06, 4.94, 1.93, 0.38, 11.5, True, NAVY, align=PP_ALIGN.CENTER, valign=MSO_ANCHOR.MIDDLE)
        if i < len(labels) - 1:
            add_text(s, "→", x + 2.09, 4.96, 0.35, 0.3, 18, True, BLUE, align=PP_ALIGN.CENTER)
        x += 2.48
    add_text(s, "La tabla temporal solo recibe y revisa la fuente; no forma parte del modelo normalizado.", 0.8, 6.05, 11.8, 0.45, 12, False, MUTED, align=PP_ALIGN.CENTER)
    add_notes(s, "El problema es que la fuente original tiene 106.555 filas y 28 columnas en una estructura plana. Eso significa que se repiten instituciones, carreras, comunas y otros valores. Primero recibimos la información en una tabla temporal de carga para revisarla y después la distribuimos en 19 tablas normalizadas. El objetivo es que EDUBIO360 pueda responder preguntas reales, por ejemplo qué institución imparte una carrera, dónde se ofrece, en qué modalidad y cuál es su costo.")


def slide3(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    set_bg(s)
    add_title(s, "2. Modelo relacional", "Separar conceptos permite reducir repetición y mantener relaciones consistentes.", 3)
    groups = [("REGION", "PROVINCIA", "COMUNA"), ("TIPO_INSTITUCION", "INSTITUCION", "OFERTA_ACADEMICA"), ("AREA_CONOCIMIENTO", "DENOMINACION_CARRERA", "CARRERA"), ("TIPO_PLAN", "PLAN_OFERTA", "MATRICULA_HISTORICA")]
    ys = [1.45, 2.55, 3.65, 4.75]
    cols = [BLUE, GREEN, GOLD, NAVY]
    for row, (a, b, c) in enumerate(groups):
        y = ys[row]
        for j, t in enumerate((a, b, c)):
            x = 0.7 + j * 3.45
            sh = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(2.7), Inches(0.62))
            sh.fill.solid()
            sh.fill.fore_color.rgb = WHITE
            sh.line.color.rgb = cols[row]
            add_text(s, t, x + 0.05, y + 0.15, 2.6, 0.28, 10.5, True, cols[row], align=PP_ALIGN.CENTER)
            if j < 2:
                add_text(s, "1:N  →", x + 2.76, y + 0.17, 0.63, 0.25, 9.5, True, MUTED, align=PP_ALIGN.CENTER)
    card(s, 10.45, 1.45, 2.28, 1.15, "19 tablas", "Modelo de negocio normalizado.", BLUE, 11.5)
    card(s, 10.45, 2.85, 2.28, 1.15, "Relaciones 1:N", "PK y FK mantienen la integridad.", GREEN, 11.5)
    card(s, 10.45, 4.25, 2.28, 1.15, "1FN · 2FN · 3FN", "Se separan conceptos y dependencias.", GOLD, 11.5)
    add_text(s, "La tabla STAGING_MATRICULA queda fuera del modelo porque solo cumple una función temporal de carga.", 0.75, 6.18, 11.8, 0.42, 11.5, False, MUTED, align=PP_ALIGN.CENTER)
    add_notes(s, "Este es el modelo relacional de EDUBIO360. Está formado por 19 tablas de negocio. La tabla STAGING_MATRICULA no se considera parte del modelo porque solo recibe temporalmente la fuente. Aplicamos normalización hasta tercera forma normal. Por ejemplo, separamos denominación de carrera, nivel, oferta académica y plan para evitar repetir información y permitir que cada concepto tenga una responsabilidad clara.")


def slide4(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    set_bg(s)
    add_title(s, "3. RECORD: una oferta como una sola estructura", "Agrupa datos relacionados y hace más claro el procesamiento de una fila compuesta.", 4)
    pill(s, "IE1.1.2 · 15%", 10.75, 0.42, 1.55, GREEN)
    card(s, 0.55, 1.45, 3.65, 2.0, "¿Por qué RECORD?", "Una oferta académica tiene carrera, nivel, institución, comuna, modalidad, jornada, matrícula y arancel. RECORD permite tratarlos como una sola estructura.", BLUE, 12.7)
    card(s, 0.55, 3.72, 3.65, 1.65, "Aporte", "Mejora la organización del bloque y evita administrar muchas variables sueltas. Su beneficio principal es claridad y coherencia.", GREEN, 12.7)
    code = "TYPE t_oferta IS RECORD (\n  nombre_carrera     ...%TYPE,\n  nombre_institucion ...%TYPE,\n  nombre_comuna      ...%TYPE,\n  valor_matricula    ...%TYPE,\n  valor_arancel      ...%TYPE\n);\n\nSELECT ... INTO v_oferta..."
    sh = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(4.55), Inches(1.45), Inches(4.08), Inches(4.45))
    sh.fill.solid()
    sh.fill.fore_color.rgb = RGBColor(15, 23, 42)
    sh.line.fill.background()
    add_text(s, "Bloque PL/SQL anónimo", 4.82, 1.72, 3.5, 0.35, 13, True, RGBColor(147, 197, 253))
    add_text(s, code, 4.82, 2.2, 3.55, 3.35, 12.2, False, RGBColor(229, 231, 235), font="Consolas")
    card(s, 8.95, 1.45, 3.8, 4.45, "v_oferta", "Carrera · Nivel\nInstitución · Comuna\nModalidad · Jornada\nMatrícula · Arancel\n\nUn solo registro lógico", GOLD, 15)
    add_notes(s, "RECORD sirve cuando varios datos pertenecen lógicamente a la misma fila. En nuestro caso una oferta tiene carrera, institución, comuna, modalidad, jornada, matrícula y arancel. En vez de mantener una variable independiente para cada dato, definimos t_oferta y trabajamos con v_oferta. No afirmamos que RECORD haga la consulta más rápida; su aporte principal es organizar el bloque y reducir variables sueltas.")


def slide5(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    set_bg(s)
    add_title(s, "4. VARRAY: conjunto pequeño y controlado", "Se utiliza cuando conocemos de antemano un límite máximo de elementos.", 5)
    pill(s, "IE1.1.2 · 15%", 10.75, 0.42, 1.55, GREEN)
    card(s, 0.55, 1.45, 3.55, 1.55, "Necesidad", "EDUBIO360 necesita recuperar un grupo reducido de alternativas. En el ejemplo se trabajan 5 ofertas con arancel positivo.", BLUE, 12.4)
    card(s, 0.55, 3.2, 3.55, 1.55, "Decisión técnica", "VARRAY(5) fija el máximo. BULK COLLECT carga varias filas desde SQL hacia PL/SQL en una sola operación.", GREEN, 12.4)
    card(s, 0.55, 4.95, 3.55, 1.25, "Cuándo no usarlo", "No es ideal si la cantidad puede crecer sin un límite conocido.", GOLD, 12.4)
    img = EVID / "02_varray.png"
    if not add_picture_if_exists(s, img, 4.45, 1.5, 8.3, 4.95):
        card(s, 4.45, 1.5, 8.3, 4.95, "Ejecución real", "TYPE t_nombres IS VARRAY(5) ...\nBULK COLLECT INTO ...\nFOR i IN 1 .. v_nombres.COUNT LOOP ...", BLUE, 18)
    add_text(s, "Ejecución real: VARRAY + BULK COLLECT + LOOP", 4.55, 6.55, 8.0, 0.32, 11, True, MUTED, align=PP_ALIGN.CENTER)
    add_notes(s, "VARRAY se justifica porque en este ejemplo conocemos el máximo: cinco alternativas. Definimos arreglos de cinco posiciones para nombres y aranceles. BULK COLLECT permite traer varias filas desde SQL hacia la colección en una operación y luego las recorremos con un LOOP. Si la cantidad fuera abierta o creciera sin límite, VARRAY dejaría de ser la mejor alternativa.")


def slide6(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    set_bg(s)
    add_title(s, "5. Cursores con y sin parámetros + loops anidados", "El objetivo es recorrer áreas y, dentro de cada una, sus ofertas académicas.", 6)
    pill(s, "IE1.2.2 · 15%", 10.75, 0.42, 1.55, GREEN)
    card(s, 0.55, 1.4, 3.0, 1.18, "Cursor sin parámetros", "c_areas recorre todas las áreas de conocimiento.", BLUE, 12.1)
    card(s, 0.55, 2.77, 3.0, 1.35, "Cursor parametrizado", "c_ofertas_por_area(p_id_area) recibe el ID del área y consulta sus ofertas.", GREEN, 12.1)
    card(s, 0.55, 4.32, 3.0, 1.35, "Loops simultáneos", "LOOP externo: áreas. LOOP interno: ofertas de cada área.", GOLD, 12.1)
    sh = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(3.85), Inches(1.4), Inches(4.15), Inches(2.0))
    sh.fill.solid()
    sh.fill.fore_color.rgb = RGBColor(15, 23, 42)
    sh.line.fill.background()
    add_text(s, "FOR a IN c_areas LOOP\n  FOR o IN c_ofertas_por_area(a.id_area) LOOP\n    DBMS_OUTPUT.PUT_LINE(...);\n  END LOOP;\nEND LOOP;", 4.15, 1.8, 3.55, 1.25, 12.2, False, RGBColor(229, 231, 235), font="Consolas")
    card(s, 3.85, 3.62, 4.15, 2.05, "¿Por qué es complejo?", "Integra varias tablas relacionadas: denominación, carrera, nivel, oferta, institución y plan. El parámetro cambia el conjunto de resultados en cada iteración.", NAVY, 12)
    img = EVID / "03_cursores_loops.png"
    if not add_picture_if_exists(s, img, 8.3, 1.4, 4.45, 4.7):
        card(s, 8.3, 1.4, 4.45, 4.7, "Ejecución real por área", "También se probó recorrido manual con OPEN · FETCH · %NOTFOUND · CLOSE.", BLUE, 17)
    add_text(s, "También se probó recorrido manual con OPEN · FETCH · %NOTFOUND · CLOSE.", 3.95, 6.15, 8.7, 0.42, 11.3, True, MUTED, align=PP_ALIGN.CENTER)
    add_notes(s, "Aquí se encuentra uno de los puntos más importantes de la pauta. c_areas es un cursor sin parámetros. c_ofertas_por_area recibe p_id_area, por lo que es parametrizado. El LOOP externo recorre áreas y el interno usa ese ID para obtener sus ofertas. Es un cursor complejo porque integra información de varias tablas: denominación, carrera, nivel, oferta académica, institución y plan. Además probamos la forma manual con OPEN, FETCH, porcentaje NOTFOUND y CLOSE.")


def slide7(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    set_bg(s)
    add_title(s, "6. Excepciones: controlar errores sin perder el flujo", "Oracle reconoce ciertos errores; las reglas propias del proyecto requieren excepciones personalizadas.", 7)
    pill(s, "IE1.3.2 · 15%", 10.75, 0.42, 1.55, GREEN)
    card(s, 0.55, 1.45, 3.35, 1.35, "NO_DATA_FOUND", "Se usa cuando SELECT INTO no devuelve ninguna fila. Ejemplo: buscar un ID de carrera inexistente.", BLUE, 12)
    card(s, 0.55, 3.0, 3.35, 1.35, "TOO_MANY_ROWS", "Se usa cuando SELECT INTO esperaba una fila pero la consulta devuelve varias.", GOLD, 12)
    card(s, 0.55, 4.55, 3.35, 1.35, "e_arancel_invalido", "Excepción definida por el usuario. Se lanza con RAISE cuando un arancel es negativo.", GREEN, 12)
    add_picture_if_exists(s, EVID / "04a_no_data_found.png", 4.2, 1.45, 4.1, 2.0)
    add_picture_if_exists(s, EVID / "04b_too_many_rows.png", 8.55, 1.45, 4.1, 2.0)
    add_picture_if_exists(s, EVID / "04c_excepcion_usuario.png", 4.2, 3.75, 4.1, 2.0)
    card(s, 8.55, 3.75, 4.1, 2.0, "Integridad y criterio", "Las excepciones controlan el flujo y entregan mensajes claros, pero no reemplazan PK, FK, NOT NULL ni CHECK.\n\nOracle = error reconocido por el motor.\nUsuario = regla específica del negocio.", NAVY, 11.3)
    add_notes(s, "Usamos dos excepciones predefinidas por Oracle y una propia. NO_DATA_FOUND aparece cuando SELECT INTO no devuelve filas. TOO_MANY_ROWS ocurre cuando esperábamos una fila y recibimos varias. e_arancel_invalido representa una regla del proyecto y la levantamos con RAISE cuando el valor es negativo. Las excepciones controlan el flujo, pero no reemplazan restricciones de integridad como PK, FK, NOT NULL o CHECK.")


def slide8(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    set_bg(s)
    add_title(s, "7. Uso futuro de objetos almacenados", "No se agregan por cantidad: cada objeto debe resolver una responsabilidad concreta.", 8)
    pill(s, "IE1.4.2 · 15%", 10.75, 0.42, 1.55, GREEN)
    items = [("PROCEDURE", "pr_generar_reporte_ofertas(p_id_area)", "Ejecutaría una tarea completa: generar ofertas para un área.", BLUE), ("FUNCTION", "fn_clasificar_arancel(p_valor_arancel)", "Recibiría un arancel y devolvería una clasificación.", GREEN), ("PACKAGE", "pkg_academico", "Agruparía procedimientos y funciones académicas relacionadas.", NAVY), ("TRIGGER", "auditoría de arancel", "Se ejecutaría ante cambios y guardaría valor anterior, nuevo valor, fecha y usuario.", GOLD)]
    xs = [0.55, 3.7, 6.85, 10.0]
    for x, (t, n, b, c) in zip(xs, items):
        card(s, x, 1.55, 2.78, 4.25, t, f"{n}\n\n{b}", c, 12)
    add_text(s, "Procedure = tarea completa   ·   Function = devuelve un valor   ·   Package = organiza lógica   ·   Trigger = responde automáticamente a eventos", 0.75, 6.18, 11.8, 0.45, 11.5, True, MUTED, align=PP_ALIGN.CENTER)
    add_notes(s, "La pauta pide explicar la utilización futura de estos objetos. No los proponemos solo por cumplir. El procedimiento generaría reportes por área, la función clasificaría un arancel y devolvería un valor, el package agruparía la lógica académica relacionada y el trigger serviría para auditoría cuando se modifiquen aranceles. Cada uno tiene una responsabilidad distinta.")


def slide9(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    set_bg(s)
    add_title(s, "8. Cómo se integrarían y qué riesgos controlaríamos", "Una solución integral necesita reutilización, dependencias claras, rendimiento y seguridad.", 9)
    pill(s, "IE1.4.2 · 15%", 10.75, 0.42, 1.55, GREEN)
    flow = [("Backend / servicio", BLUE), ("pkg_academico", NAVY), ("Procedure", GREEN), ("Function", GREEN), ("Tablas Oracle", BLUE), ("Trigger auditoría", GOLD)]
    x = 0.55
    for i, (t, c) in enumerate(flow):
        sh = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(1.55), Inches(1.82), Inches(0.72))
        sh.fill.solid()
        sh.fill.fore_color.rgb = WHITE
        sh.line.color.rgb = c
        add_text(s, t, x + 0.03, 1.75, 1.76, 0.25, 10.5, True, c, align=PP_ALIGN.CENTER)
        if i < len(flow) - 1:
            add_text(s, "→", x + 1.84, 1.76, 0.3, 0.25, 14, True, MUTED, align=PP_ALIGN.CENTER)
        x += 2.08
    card(s, 0.55, 3.0, 2.95, 2.25, "Rendimiento", "Evitar lógica pesada y procesamiento fila por fila cuando SQL pueda resolverlo mejor.", BLUE, 12)
    card(s, 3.72, 3.0, 2.95, 2.25, "Mantenimiento", "Una responsabilidad clara por objeto y revisión de dependencias cuando cambian tablas o columnas.", GREEN, 12)
    card(s, 6.89, 3.0, 2.95, 2.25, "Seguridad", "Permisos solo para usuarios o procesos que realmente necesiten ejecutar o modificar los objetos.", NAVY, 12)
    card(s, 10.06, 3.0, 2.72, 2.25, "Integridad", "Los triggers apoyan auditoría, pero no reemplazan PK, FK, NOT NULL ni CHECK.", GOLD, 12)
    add_text(s, "La estrategia busca reutilización sin esconder toda la lógica del sistema dentro de Oracle.", 0.75, 6.15, 11.8, 0.45, 12, True, MUTED, align=PP_ALIGN.CENTER)
    add_notes(s, "Estos objetos funcionarían en conjunto, pero hay que controlar sus riesgos. En rendimiento evitaríamos procesos fila por fila si SQL puede resolverlos mejor. En mantenimiento cada objeto debe tener una responsabilidad clara y hay que revisar dependencias si cambia una tabla. En seguridad no todos los usuarios deben poder ejecutarlos o modificarlos. Y los triggers sirven como apoyo de auditoría, no como reemplazo de las restricciones de integridad.")


def slide10(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    set_bg(s, WHITE)
    add_title(s, "9. Conclusión", "EDUBIO360 transforma una fuente plana en información estructurada, procesable y reutilizable.", 10)
    card(s, 0.7, 1.55, 3.75, 3.55, "1 · Datos organizados", "106.555 registros se cargan, revisan y distribuyen en 19 tablas normalizadas.", BLUE, 16)
    card(s, 4.79, 1.55, 3.75, 3.55, "2 · PL/SQL con propósito", "RECORD agrupa, VARRAY limita colecciones, cursores recorren y excepciones controlan situaciones reales.", GREEN, 15)
    card(s, 8.88, 1.55, 3.75, 3.55, "3 · Base para crecer", "Procedures, Functions, Packages y Triggers se incorporarían cuando exista una necesidad concreta y justificable.", GOLD, 14.7)
    add_text(s, "Resultado: una base de datos consistente que puede ser consumida por el backend sin depender de la planilla original.", 1.0, 5.65, 11.3, 0.7, 17, True, NAVY, align=PP_ALIGN.CENTER)
    add_text(s, "Gracias", 5.55, 6.55, 2.2, 0.4, 18, True, BLUE, align=PP_ALIGN.CENTER)
    add_notes(s, "En conclusión, primero transformamos una fuente plana de 106.555 registros en 19 tablas normalizadas. Después utilizamos PL/SQL con una finalidad concreta: RECORD para agrupar datos, VARRAY para conjuntos pequeños, cursores para recorridos relacionados y excepciones para controlar errores. Finalmente dejamos definida una estrategia futura para procedures, functions, packages y triggers. Con esto EDUBIO360 queda preparado para que otros componentes consulten información ordenada y consistente.")


def normalize_zip(path: Path):
    with zipfile.ZipFile(path, "r") as zin:
        entries = [(info.filename, zin.read(info.filename), info.external_attr) for info in zin.infolist()]
    tmp = path.with_suffix(".tmp")
    with zipfile.ZipFile(tmp, "w", compression=zipfile.ZIP_DEFLATED, compresslevel=9) as zout:
        for name, data, attrs in sorted(entries, key=lambda x: x[0]):
            zi = zipfile.ZipInfo(name, (1980, 1, 1, 0, 0, 0))
            zi.compress_type = zipfile.ZIP_DEFLATED
            zi.external_attr = attrs
            zout.writestr(zi, data)
    tmp.replace(path)


def main():
    OUT.parent.mkdir(parents=True, exist_ok=True)
    prs = Presentation()
    prs.slide_width = W
    prs.slide_height = H
    prs.core_properties.title = "EDUBIO360 - Evaluación Parcial N° 1"
    prs.core_properties.subject = "Presentación BDY1103"
    prs.core_properties.author = "Javier Pérez y Nicolás Vega"
    prs.core_properties.created = datetime(2026, 9, 23, 12, 0, 0)
    prs.core_properties.modified = datetime(2026, 9, 23, 12, 0, 0)
    for fn in (slide1, slide2, slide3, slide4, slide5, slide6, slide7, slide8, slide9, slide10):
        fn(prs)
    prs.save(OUT)
    normalize_zip(OUT)
    print(f"Presentación generada: {OUT}")


if __name__ == "__main__":
    main()
